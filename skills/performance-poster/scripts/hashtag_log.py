#!/usr/bin/env python3
"""Hashtag log: which hashtags went on which post, and how each post did, so the tag sets can be judged on
Pecuna's own results instead of vendor studies.

The log is a CSV (default: references/hashtags/hashtag-log.csv). Like the headline log, the skill's copy doesn't
survive between claude.ai sessions: keep the exported .xlsx as the master and `import` it first when Taylor attaches it.

Usage:
  python3 hashtag_log.py add --video "AI Harnesses" --platform linkedin --post "native cut" \
      --tags "#ClaudeCode #AIAgents #AIAutomation" [--pillar "AI agents"] [--date 2026-10-22] [--notes "..."]
  python3 hashtag_log.py update --id 3 [--reach 1840] [--impressions 2210] [--engagement 4.2] [--notes "..."]
  python3 hashtag_log.py list [--platform instagram]
  python3 hashtag_log.py report [--platform linkedin]   # per hashtag: posts, median reach, median engagement
  python3 hashtag_log.py export hashtags.xlsx           # needs openpyxl
  python3 hashtag_log.py import hashtags.xlsx

platform: youtube | linkedin | instagram. Log results at 48 hours (reach/impressions) and 30 days (engagement).
A report is correlation on a small sample: treat it as a hint about which tags to re-test, never as proof.
"""
import argparse
import csv
import datetime as dt
import os
import re
import statistics
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_LOG = os.path.join(HERE, '..', 'references', 'hashtags', 'hashtag-log.csv')
COLS = ['id', 'date', 'video', 'platform', 'post', 'pillar', 'hashtags', 'count', 'reach', 'impressions', 'engagement', 'notes']
PLATFORMS = ('youtube', 'linkedin', 'instagram')


def load(path):
    if not os.path.exists(path):
        return []
    with open(path, newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def save(path, rows):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, '') for c in COLS})


def parse_tags(s):
    return re.findall(r'#\w+', s or '')


def num(v):
    try:
        return float(str(v).replace(',', '').replace('%', '').strip())
    except (TypeError, ValueError):
        return None


def report(rows, platform=None):
    """Per hashtag (case-insensitive): number of posts, median reach and median engagement where logged."""
    by = {}
    for r in rows:
        if platform and r['platform'] != platform:
            continue
        for t in parse_tags(r['hashtags']):
            d = by.setdefault(t.lower(), {'tag': t, 'posts': 0, 'reach': [], 'engagement': []})
            d['posts'] += 1
            for k in ('reach', 'engagement'):
                if num(r.get(k)) is not None:
                    d[k].append(num(r[k]))
    out = []
    for d in by.values():
        out.append({'tag': d['tag'], 'posts': d['posts'],
                    'median_reach': statistics.median(d['reach']) if d['reach'] else None,
                    'median_engagement': statistics.median(d['engagement']) if d['engagement'] else None})
    return sorted(out, key=lambda x: (-(x['median_reach'] or -1), -x['posts']))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--log', default=DEFAULT_LOG)
    sub = ap.add_subparsers(dest='cmd', required=True)
    a = sub.add_parser('add')
    for k in ('video', 'post', 'pillar', 'date', 'notes'):
        a.add_argument('--' + k, default='')
    a.add_argument('--platform', required=True, choices=PLATFORMS)
    a.add_argument('--tags', required=True)
    u = sub.add_parser('update')
    u.add_argument('--id', required=True)
    for k in ('reach', 'impressions', 'engagement', 'notes'):
        u.add_argument('--' + k)
    for name in ('list', 'report'):
        s = sub.add_parser(name)
        s.add_argument('--platform', choices=PLATFORMS)
    e = sub.add_parser('export')
    e.add_argument('xlsx')
    i = sub.add_parser('import')
    i.add_argument('xlsx')
    args = ap.parse_args(argv)
    rows = load(args.log)

    if args.cmd == 'add':
        tags = parse_tags(args.tags)
        if not tags:
            sys.exit('no hashtags found in --tags')
        nid = max([int(r['id']) for r in rows if r['id'].isdigit()] or [0]) + 1
        rows.append({'id': nid, 'date': args.date or dt.date.today().isoformat(), 'video': args.video,
                     'platform': args.platform, 'post': args.post, 'pillar': args.pillar,
                     'hashtags': ' '.join(tags), 'count': len(tags), 'notes': args.notes})
        save(args.log, rows)
        print(f'added #{nid}: {args.platform} · {" ".join(tags)}')
    elif args.cmd == 'update':
        hit = [r for r in rows if r['id'] == str(args.id)]
        if not hit:
            sys.exit(f'no row with id {args.id}')
        for k in ('reach', 'impressions', 'engagement', 'notes'):
            v = getattr(args, k)
            if v is not None:
                hit[0][k] = v
        save(args.log, rows)
        print(f'updated #{args.id}')
    elif args.cmd == 'list':
        for r in rows:
            if not args.platform or r['platform'] == args.platform:
                print(f"#{r['id']} {r['date']} {r['platform']:9} {r['video']} · {r['hashtags']} · reach {r['reach'] or '-'}")
    elif args.cmd == 'report':
        rep = report(rows, args.platform)
        if not rep:
            print('nothing logged yet')
        for d in rep:
            print(f"{d['tag']:24} posts {d['posts']:3}  median reach {d['median_reach'] if d['median_reach'] is not None else '-':>8}"
                  f"  median engagement {d['median_engagement'] if d['median_engagement'] is not None else '-'}")
        print('Correlation on a small sample: re-test before changing a tag set.')
    elif args.cmd == 'export':
        try:
            from openpyxl import Workbook
        except ImportError:
            sys.exit('export needs openpyxl (pip install openpyxl)')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Hashtag posts'
        ws.append(COLS)
        for r in rows:
            ws.append([r.get(c, '') for c in COLS])
        ws2 = wb.create_sheet('By hashtag')
        ws2.append(['hashtag', 'posts', 'median reach', 'median engagement'])
        for d in report(rows):
            ws2.append([d['tag'], d['posts'], d['median_reach'], d['median_engagement']])
        wb.save(args.xlsx)
        print(f'wrote {args.xlsx} ({len(rows)} posts)')
    elif args.cmd == 'import':
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit('import needs openpyxl (pip install openpyxl)')
        ws = load_workbook(args.xlsx, read_only=True)['Hashtag posts']
        it = ws.iter_rows(values_only=True)
        head = [str(h) for h in next(it)]
        new = [{h: ('' if v is None else str(v)) for h, v in zip(head, row)} for row in it if any(v is not None for v in row)]
        save(args.log, new)
        print(f'imported {len(new)} posts into {args.log}')


if __name__ == '__main__':
    main()
